# Fixture com grupo aninhado e travas, para testar unlock.py
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])

# tres shapes que serao agrupados manualmente no XML
coords = [(1,1,2,1.5), (3.5,1,2,1.5), (6,1,2,1.5)]
els = []
for (x,y,w,h) in coords:
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = RGBColor(0x0E,0x7C,0x74)
    sp.text_frame.text = f"bloco {x}"
    els.append(sp._element)

spTree = s.shapes._spTree
grp_xml = f'''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:nvGrpSpPr><p:cNvPr id="900" name="grupo externo"/><p:cNvGrpSpPr>
   <a:grpSpLocks noUngrp="1" noSelect="1"/></p:cNvGrpSpPr><p:nvPr/></p:nvGrpSpPr>
 <p:grpSpPr><a:xfrm>
   <a:off x="{Emu(Inches(1))}" y="{Emu(Inches(1))}"/>
   <a:ext cx="{Emu(Inches(7))}" cy="{Emu(Inches(1.5))}"/>
   <a:chOff x="{Emu(Inches(1))}" y="{Emu(Inches(1))}"/>
   <a:chExt cx="{Emu(Inches(7))}" cy="{Emu(Inches(1.5))}"/>
 </a:xfrm></p:grpSpPr>
</p:grpSp>'''
grp = etree.fromstring(grp_xml)

# grupo interno com os dois primeiros
inner_xml = f'''<p:grpSp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
 <p:nvGrpSpPr><p:cNvPr id="901" name="grupo interno"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
 <p:grpSpPr><a:xfrm>
   <a:off x="{Emu(Inches(1))}" y="{Emu(Inches(1))}"/>
   <a:ext cx="{Emu(Inches(4.5))}" cy="{Emu(Inches(1.5))}"/>
   <a:chOff x="{Emu(Inches(1))}" y="{Emu(Inches(1))}"/>
   <a:chExt cx="{Emu(Inches(4.5))}" cy="{Emu(Inches(1.5))}"/>
 </a:xfrm></p:grpSpPr>
</p:grpSp>'''
inner = etree.fromstring(inner_xml)

for e in els[:2]:
    spTree.remove(e); inner.append(e)
grp.append(inner)
spTree.remove(els[2]); grp.append(els[2])

# trava tambem em um shape solto
sp4 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(1), Inches(2), Inches(1.5))
locks = etree.SubElement(sp4._element.find(qn('p:nvSpPr')).find(qn('p:cNvSpPr')), qn('a:spLocks'))
locks.set('noMove','1'); locks.set('noResize','1'); locks.set('noTextEdit','1')

spTree.append(grp)
prs.save("fix_grupo.pptx")
print("fixture com grupo aninhado criada")
