#!/usr/bin/env python3
"""
audit_pptx.py - Modulo de triagem para conversao de apresentacoes em editaveis.

Classifica um .pptx e mede, por slide, quanto do conteudo ja e nativo e quanto
esta preso em raster. E o modulo 1 do pipeline: define a rota de conversao e
estabelece a linha de base contra a qual o resultado sera medido.

Uso:
    python3 audit_pptx.py arquivo.pptx [--json saida.json]

Dependencias: python-pptx, lxml (stdlib para o resto).
"""

import argparse
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Limiares. Documentados aqui porque sao decisoes de projeto, nao constantes
# naturais: ajuste com base no seu corpus antes de usar em producao.
COBERTURA_ACHATADO = 0.85   # imagem cobrindo >=85% do slide sugere slide achatado
CHARS_MINIMOS = 15          # abaixo disso, o slide nao tem texto util nativo
COBERTURA_MISTO = 0.35      # acima disso ja ha raster relevante competindo


def _area_emu(prs):
    return int(prs.slide_width) * int(prs.slide_height)


def _shape_area(shape):
    try:
        w = int(shape.width or 0)
        h = int(shape.height or 0)
    except (TypeError, ValueError):
        return 0
    return max(w, 0) * max(h, 0)


def _walk(shapes):
    """Percorre shapes recursivamente, entrando em grupos. Devolve (shape, profundidade)."""
    for sh in shapes:
        yield sh, 0
        if sh.shape_type is not None and getattr(sh, "shapes", None) is not None:
            for sub, d in _walk(sh.shapes):
                yield sub, d + 1


def _tem_travas(shape):
    """Detecta spLocks/grpSpLocks restritivos no XML do shape."""
    try:
        xml = shape._element.xml
    except Exception:
        return []
    travas = []
    for attr in ("noSelect", "noMove", "noResize", "noEditPoints",
                 "noTextEdit", "noUngrp", "noChangeShapeType"):
        if re.search(rf'{attr}="(1|true)"', xml):
            travas.append(attr)
    return travas


def paleta_do_tema(caminho):
    """Le os srgbClr declarados no tema e nos slides, por frequencia."""
    tema, slides = Counter(), Counter()
    with zipfile.ZipFile(caminho) as z:
        for nome in z.namelist():
            if nome.startswith("ppt/theme/") and nome.endswith(".xml"):
                alvo = tema
            elif nome.startswith("ppt/slides/slide") and nome.endswith(".xml"):
                alvo = slides
            else:
                continue
            texto = z.read(nome).decode("utf-8", "ignore")
            for hexa in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', texto):
                alvo["#" + hexa.upper()] += 1
    return tema, slides


def auditar(caminho):
    prs = Presentation(caminho)
    area_slide = _area_emu(prs)
    largura_pol = Emu(prs.slide_width).inches
    altura_pol = Emu(prs.slide_height).inches

    relatorio = {
        "arquivo": Path(caminho).name,
        "palco_polegadas": [round(largura_pol, 3), round(altura_pol, 3)],
        "n_slides": len(prs.slides),
        "slides": [],
    }

    for i, slide in enumerate(prs.slides, 1):
        cont = Counter()
        chars = 0
        area_raster = 0
        maior_imagem = 0.0
        travas = []
        profundidade_max = 0

        for sh, prof in _walk(slide.shapes):
            profundidade_max = max(profundidade_max, prof)
            tipo = str(sh.shape_type)

            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    cont["caixas_texto"] += 1
                    chars += len(t)

            if sh.shape_type is not None and "PICTURE" in tipo:
                cont["imagens"] += 1
                a = _shape_area(sh)
                area_raster += a
                if area_slide:
                    maior_imagem = max(maior_imagem, a / area_slide)
            elif getattr(sh, "has_table", False) and sh.has_table:
                cont["tabelas"] += 1
            elif getattr(sh, "has_chart", False) and sh.has_chart:
                cont["graficos"] += 1
            elif sh.shape_type is not None and "GROUP" in tipo:
                cont["grupos"] += 1
            elif sh.shape_type is not None and "AUTO_SHAPE" in tipo:
                cont["formas"] += 1
            elif sh.shape_type is not None and "FREEFORM" in tipo:
                cont["formas_livres"] += 1

            t_sh = _tem_travas(sh)
            if t_sh:
                travas.append({"shape": sh.shape_id, "travas": t_sh})

        cobertura = min(area_raster / area_slide, 1.0) if area_slide else 0.0

        # Classificacao da rota
        if maior_imagem >= COBERTURA_ACHATADO and chars < CHARS_MINIMOS:
            classe, confianca = "ACHATADO", "alta"
        elif cobertura >= COBERTURA_ACHATADO and chars < CHARS_MINIMOS:
            classe, confianca = "ACHATADO", "media"
        elif cobertura >= COBERTURA_MISTO:
            classe, confianca = "MISTO", "media"
        else:
            classe, confianca = "NATIVO", "alta"

        # Indice de Editabilidade: fracao da area do slide NAO ocupada por raster,
        # com penalidade quando nao ha texto nativo algum.
        ie = 1.0 - cobertura
        if chars < CHARS_MINIMOS and cont["imagens"] > 0:
            ie *= 0.5

        relatorio["slides"].append({
            "slide": i,
            "classe": classe,
            "confianca": confianca,
            "indice_editabilidade": round(ie, 3),
            "cobertura_raster": round(cobertura, 3),
            "maior_imagem_frac": round(maior_imagem, 3),
            "caracteres_nativos": chars,
            "objetos": dict(cont),
            "profundidade_agrupamento": profundidade_max,
            "travas": travas,
        })

    tema, slides_cores = paleta_do_tema(caminho)
    relatorio["paleta_tema"] = tema.most_common(16)
    relatorio["paleta_slides"] = slides_cores.most_common(16)

    ies = [s["indice_editabilidade"] for s in relatorio["slides"]]
    relatorio["indice_editabilidade_global"] = round(sum(ies) / len(ies), 3) if ies else 0.0
    relatorio["distribuicao_classes"] = dict(
        Counter(s["classe"] for s in relatorio["slides"]))
    relatorio["rota_sugerida"] = _rota(relatorio["distribuicao_classes"])
    return relatorio


def _rota(dist):
    achatado = dist.get("ACHATADO", 0)
    misto = dist.get("MISTO", 0)
    total = sum(dist.values()) or 1
    if achatado / total >= 0.5:
        return ("TRILHA C (raster): OCR + analise de layout + vetorizacao. "
                "Automacao parcial, revisao humana obrigatoria.")
    if (achatado + misto) / total >= 0.3:
        return ("TRILHA B (misto): reconstruir apenas os slides marcados, "
                "preservar os nativos. Automacao alta nos nativos.")
    return ("TRILHA A (nativo): destravar, desagrupar e promover elementos de "
            "layout/master. Automacao total, deterministica.")


def imprimir(r):
    print(f"\nArquivo: {r['arquivo']}")
    print(f"Palco:   {r['palco_polegadas'][0]}\" x {r['palco_polegadas'][1]}\"   "
          f"Slides: {r['n_slides']}")
    print(f"Indice de Editabilidade global: {r['indice_editabilidade_global']}")
    print(f"Classes: {r['distribuicao_classes']}")
    print(f"Rota:    {r['rota_sugerida']}\n")

    cab = f"{'#':>3}  {'classe':<9} {'IE':>5} {'raster':>7} {'chars':>6}  objetos"
    print(cab)
    print("-" * max(len(cab), 78))
    for s in r["slides"]:
        obj = ", ".join(f"{k}={v}" for k, v in sorted(s["objetos"].items())) or "-"
        print(f"{s['slide']:>3}  {s['classe']:<9} {s['indice_editabilidade']:>5} "
              f"{s['cobertura_raster']:>7} {s['caracteres_nativos']:>6}  {obj}")
        if s["travas"]:
            print(f"     travas: {s['travas']}")

    print("\nPaleta declarada no tema (hex, frequencia):")
    for h, n in r["paleta_tema"][:10]:
        print(f"  {h}  {n}")
    if r["paleta_slides"]:
        print("Cores literais nos slides (fora do tema, menos editaveis):")
        for h, n in r["paleta_slides"][:10]:
            print(f"  {h}  {n}")


def main():
    ap = argparse.ArgumentParser(description="Triagem de editabilidade de .pptx")
    ap.add_argument("arquivo")
    ap.add_argument("--json", help="grava o relatorio em JSON")
    a = ap.parse_args()

    if not Path(a.arquivo).exists():
        sys.exit(f"erro: arquivo nao encontrado: {a.arquivo}")
    if not a.arquivo.lower().endswith((".pptx", ".potx")):
        sys.exit("erro: entrada deve ser .pptx ou .potx. "
                 "Converta .ppt antes: soffice --headless --convert-to pptx")

    r = auditar(a.arquivo)
    imprimir(r)
    if a.json:
        Path(a.json).write_text(json.dumps(r, ensure_ascii=False, indent=2),
                                encoding="utf-8")
        print(f"\nJSON gravado em {a.json}")


if __name__ == "__main__":
    main()
